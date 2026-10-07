"""Products tab for webui.py: edit data/products.json and attach images to each product.

Images are saved to data/product_images/ and stored in each product as
    "images": ["data/product_images/P001_1700000000.jpg", ...]

Add flow (all from the web UI, no browser extension):
  * paste a TikTok product link + Add -> the webui finds Chrome (or Edge) on this machine, starts it
    with a separate profile (data/chrome_profile, remote debugging on localhost only), opens the page,
    reads name/price/images and closes the tab. If TikTok shows a Security Check, solve it in that
    window. Set the CHROME_PATH environment variable if the browser can't be found automatically.
  * paste JSON copied by the bookmarklet + Add -> adds it directly.
  * empty + Add -> blank product.
"""
import csv
import io
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import time
import unicodedata
from pathlib import Path

import requests
from nicegui import ui, app, run

PRODUCTS_PATH = Path("data/products.json")
IMG_DIR = Path("data/product_images")
IMG_URL = "/product_images"
IMG_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif"}

CDP_PORT = 9222
CDP_URL = f"http://127.0.0.1:{CDP_PORT}"
CHROME_PROFILE = Path("data/chrome_profile")  # separate profile, lives inside the project

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")

TEXT_FIELDS = [
    ("name", "Name"), ("price", "Price"), ("original_price", "Original price"),
    ("description", "Description"), ("sizes_colors", "Sizes / colors"),
    ("how_to_use", "How to use"), ("shipping", "Shipping"),
    ("return_policy", "Return policy"), ("stock_note", "Stock note"),
]


# Filled into every NEW product (edit the text here to change the defaults).
NEW_PRODUCT_DEFAULTS = {
    "shipping": "Shop gửi hàng qua TikTok Shop, thời gian giao tuỳ khu vực.",
    "return_policy": "Được đổi trả theo chính sách của TikTok Shop.",
    "stock_note": "Số lượng có hạn trong phiên live này.",
}


def _load() -> dict:
    data = {"shop_name": "", "products": []}
    if PRODUCTS_PATH.exists():
        with open(PRODUCTS_PATH, encoding="utf-8") as f:
            data = json.load(f)
    data.setdefault("shop_name", "")
    data.setdefault("products", [])
    for p in data["products"]:
        _normalize(p)
    return data


def _normalize(p: dict) -> dict:
    for key, _ in TEXT_FIELDS:
        p.setdefault(key, "")
    p.setdefault("id", "")
    p.setdefault("order", 0)
    p.setdefault("active", True)
    p.setdefault("aliases", [])
    p.setdefault("highlights", [])
    p.setdefault("faq", [])
    p.setdefault("images", [])
    p.setdefault("intro", "")
    return p


def _save(data: dict) -> None:
    PRODUCTS_PATH.parent.mkdir(parents=True, exist_ok=True)
    tmp = PRODUCTS_PATH.with_suffix(".json.tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    os.replace(tmp, PRODUCTS_PATH)


# list <-> textarea converters (used by bind_value)
def _to_lines(v):
    return "\n".join(v or [])


def _from_lines(t):
    return [x.strip() for x in (t or "").splitlines() if x.strip()]


def _faq_to_text(v):
    return "\n".join(f"{i.get('q', '')} | {i.get('a', '')}" for i in (v or []))


def _faq_from_text(t):
    out = []
    for line in (t or "").splitlines():
        if "|" in line:
            q, a = line.split("|", 1)
            if q.strip() and a.strip():
                out.append({"q": q.strip(), "a": a.strip()})
    return out


# ---------- import from a TikTok Shop export (.xlsx / .csv) ----------
# header (accent/case-insensitive) -> product field. Add aliases here if your export differs.
HEADER_ALIASES = {
    "tiktok_id": ["product id", "id san pham", "ma san pham", "item id"],
    "name": ["product name", "ten san pham", "name", "title", "ten"],
    "price": ["price", "gia", "gia ban", "retail price", "sale price", "gia khuyen mai"],
    "original_price": ["original price", "gia goc", "list price", "gia niem yet"],
    "description": ["description", "product description", "mo ta", "mo ta san pham"],
    "sizes_colors": ["variation", "variations", "variant", "phan loai", "ten phan loai", "sku name"],
}


def _fold(t) -> str:
    t = unicodedata.normalize("NFD", str(t or "").lower().replace("đ", "d"))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


def _map_header(row) -> dict:
    """column index -> field, exact alias match first, then 'alias contained in header'."""
    mapping = {}
    for i, h in enumerate(row):
        f = _fold(h)
        if not f:
            continue
        hit = next((k for k, al in HEADER_ALIASES.items() if f in al and k not in mapping.values()), None)
        if not hit:
            hit = next((k for k, al in HEADER_ALIASES.items()
                        if k not in mapping.values() and any(a in f for a in al)), None)
        if hit:
            mapping[i] = hit
    return mapping


def _money(v) -> str:
    s = str(v if v is not None else "").strip()
    if not s:
        return ""
    try:
        n = float(s.replace(",", "")) if re.fullmatch(r"[\d,]+(\.\d+)?", s) else None
    except ValueError:
        n = None
    return f"{int(n):,}".replace(",", ".") + "đ" if n is not None else s


def _read_table(name: str, raw: bytes):
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook  # pip install openpyxl
        ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    for enc in ("utf-8-sig", "cp1258", "latin-1"):
        try:
            return list(csv.reader(io.StringIO(raw.decode(enc))))
        except UnicodeDecodeError:
            continue
    return []


def import_rows(data: dict, rows) -> tuple:
    """Merge table rows into data['products']. Returns (added, updated, headers_seen).
    Not wired to a button right now; kept for bulk import from a Seller Center export."""
    best = max(range(min(15, len(rows))), key=lambda i: len(_map_header(rows[i])), default=None)
    mapping = _map_header(rows[best]) if best is not None else {}
    if "name" not in mapping.values():
        return 0, 0, [str(c) for c in (rows[best] if best is not None else []) if c]
    groups = {}  # one product may span several variation rows
    for r in rows[best + 1:]:
        item = {mapping[i]: str(r[i]).strip() for i in mapping if i < len(r) and r[i] not in (None, "")}
        if not item.get("name"):
            continue
        g = groups.setdefault(item.get("tiktok_id") or _fold(item["name"]), item)
        if g is not item and item.get("sizes_colors"):
            old = g.get("sizes_colors", "")
            if item["sizes_colors"] not in old.split(", "):
                g["sizes_colors"] = ", ".join(x for x in (old, item["sizes_colors"]) if x)
    added = updated = 0
    for key, it in groups.items():
        cur = next((p for p in data["products"]
                    if (it.get("tiktok_id") and p.get("tiktok_id") == it["tiktok_id"])
                    or _fold(p["name"]) == _fold(it["name"])), None)
        for k in ("price", "original_price"):
            if k in it:
                it[k] = _money(it[k])
        if cur:
            for k, v in it.items():  # prices/name always refresh; text only fills blanks
                if k in ("name", "price", "original_price", "tiktok_id") or not cur.get(k):
                    cur[k] = v
            updated += 1
        else:
            n = len(data["products"]) + 1
            ids = {p["id"] for p in data["products"]}
            while f"P{n:03d}" in ids:
                n += 1
            order = max([p["order"] for p in data["products"]] + [0]) + 1
            data["products"].append(_normalize({**NEW_PRODUCT_DEFAULTS, **it, "id": f"P{n:03d}", "order": order}))
            added += 1
    return added, updated, []


# ---------- read a product page through the user's own Chrome ----------
_PAGE_JS = r"""() => {
  const good = s => /tiktokcdn|ibyteimg|byteimg/.test(s) && !/avatar|logo|icon|emoji/.test(s);
  const imgs = [...document.images].filter(i => i.naturalWidth >= 200)
      .map(i => i.currentSrc || i.src).filter(good);
  const body = document.body.innerText || "";
  return {
    title: document.title,
    h1: (document.querySelector("h1") || {}).innerText || "",
    prices: body.match(/[\d.,]+\s*[₫đ]/g) || [],
    imgs, body: body.slice(0, 8000),
  };
}"""


_PRICE_RE = re.compile(r"₫\s*\d[\d.,]*|\d[\d.,]*\s*(?:₫|đ(?!\w))")


def _prices(body: str) -> list:
    return list(dict.fromkeys(m.strip() for m in _PRICE_RE.findall(body or "")))


def _norm_price(s: str) -> str:
    d = re.sub(r"[^\d.,]", "", s or "")
    return d + "₫" if d else ""


_OPT_LABEL = re.compile(r"^(Màu sắc|Màu|Kích cỡ|Kích thước|Size|Phân loại|Kiểu)\s*:\s*(.*)$", re.I)
_OPT_STOP = re.compile(r"^(Số lượng|Quantity|Mua ngay|Thêm vào giỏ|Buy now|Add to cart)", re.I)


def _split_title(title: str):
    """Marketing titles are long: short name = first clause, full title kept as the description."""
    t = " ".join((title or "").split())
    m = re.search(r"[,.;|]| - ", t[20:])
    name = t[:20 + m.start()].strip() if m else t
    if len(name) > 70:
        name = name[:70].rsplit(" ", 1)[0]
    return name, (t if name != t else "")


def _parse_options(body: str) -> str:
    """'Màu sắc: A, B; Kích cỡ: S, M, L' read from the rendered page text (labels end with ':')."""
    groups, cur = [], None
    for line in (ln.strip() for ln in body.splitlines()):
        if not line:
            continue
        if groups and _OPT_STOP.match(line):
            break  # ignore anything below the option block (reviews etc.)
        m = _OPT_LABEL.match(line)
        if m:
            cur = [m.group(1).strip(), []]
            groups.append(cur)
            if m.group(2).strip():
                cur[1].append(m.group(2).strip())  # selected value shown on the label line
            continue
        if cur is not None and len(line) <= 40 and "₫" not in line and line not in cur[1] and len(cur[1]) < 20:
            cur[1].append(line)
    return "; ".join(f"{k}: {', '.join(v)}" for k, v in groups if v)


def _find_chrome():
    """Locate Chrome (or Edge as a fallback) on Windows / macOS / Linux. None if not found."""
    env = os.environ.get("CHROME_PATH")
    if env and Path(env).exists():
        return env
    if sys.platform == "win32":
        import winreg
        for exe in ("chrome.exe", "msedge.exe"):
            for root in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
                try:
                    with winreg.OpenKey(root, rf"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\{exe}") as k:
                        p = winreg.QueryValue(k, None)
                        if p and Path(p).exists():
                            return p
                except OSError:
                    pass
        for base in (os.environ.get("PROGRAMFILES"), os.environ.get("PROGRAMFILES(X86)"),
                     os.environ.get("LOCALAPPDATA")):
            for sub in (r"Google\Chrome\Application\chrome.exe", r"Microsoft\Edge\Application\msedge.exe"):
                if base and (Path(base) / sub).exists():
                    return str(Path(base) / sub)
    elif sys.platform == "darwin":
        for p in ("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
                  "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge"):
            if Path(p).exists():
                return p
    for name in ("google-chrome", "google-chrome-stable", "chromium", "chromium-browser",
                 "chrome", "microsoft-edge"):
        p = shutil.which(name)
        if p:
            return p
    return None


def _cdp_up() -> bool:
    try:
        with socket.create_connection(("127.0.0.1", CDP_PORT), timeout=1):
            return True
    except OSError:
        return False


def _ensure_chrome() -> None:
    """Start the helper browser (separate profile, debug port on localhost) if it isn't running."""
    if _cdp_up():
        return
    exe = _find_chrome()
    if not exe:
        raise RuntimeError("Chrome/Edge not found. Set the CHROME_PATH environment variable to its full path.")
    first_run = not CHROME_PROFILE.exists()
    args = [exe, f"--remote-debugging-port={CDP_PORT}",
            f"--user-data-dir={CHROME_PROFILE.resolve()}", "--no-first-run"]
    if first_run:
        args.append("https://www.tiktok.com")  # log in / pass the check once in this profile
    subprocess.Popen(args)
    for _ in range(40):
        time.sleep(0.5)
        if _cdp_up():
            return
    raise RuntimeError("Started the browser but its debug port didn't open. "
                       "Close any window using data/chrome_profile and try again.")


def fetch_tiktok_product(url: str) -> dict:
    """Starts the helper browser if needed, opens a tab, reads it, closes it.
    If a Security Check shows up, solve it in that window (waits up to ~2 min)."""
    from playwright.sync_api import sync_playwright  # pip install playwright (no browser download needed)
    _ensure_chrome()
    r, blocked = None, True
    with sync_playwright() as pw:
        browser = pw.chromium.connect_over_cdp(CDP_URL)
        page = browser.contexts[0].new_page()
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=45000)
            page.bring_to_front()
            for tries in range(60):
                page.wait_for_timeout(2000)
                try:
                    r = page.evaluate(_PAGE_JS)
                except Exception:
                    continue  # page is still redirecting
                blocked = bool(re.search(r"Security Check|Xác minh|captcha", r["body"], re.I))
                if not blocked and r["h1"] and (_prices(r["body"]) or tries >= 6):
                    break
        finally:
            page.close()
    if not r or blocked:
        return {}

    Path("data").mkdir(exist_ok=True)  # debug dump: send this if fields come out wrong
    Path("data/debug_tiktok.txt").write_text(json.dumps(r, ensure_ascii=False, indent=2), encoding="utf-8")

    full_title = r["h1"] or (r["title"] if r["title"] != "TikTok Shop" else "")
    name, description = _split_title(full_title)
    prices = list(dict.fromkeys(_norm_price(p) for p in _prices(r["body"])))
    num = lambda s: int(re.sub(r"\D", "", s) or 0)
    seen, imgs = set(), []
    for u in r["imgs"]:
        k = u.split("?")[0].split("~")[0]
        if k not in seen:
            seen.add(k)
            imgs.append(u)
    return {
        "name": name,
        "description": description,
        "sizes_colors": _parse_options(r["body"]),
        "price": prices[0] if prices else "",
        "original_price": (max(prices[:3], key=num) if len(prices) > 1 and num(max(prices[:3], key=num)) > num(prices[0])
                           else ""),
        "images": imgs[:6],
    }


def _download_images(pid: str, urls) -> list:
    out = []
    for u in urls:
        try:
            resp = requests.get(u, headers={"User-Agent": UA, "Referer": "https://shop.tiktok.com/"}, timeout=20)
            ext = {"image/png": ".png", "image/webp": ".webp", "image/gif": ".gif"}.get(
                resp.headers.get("content-type", "").split(";")[0], ".jpg")
            fname = f"{pid}_{int(time.time() * 1000)}{ext}"
            (IMG_DIR / fname).write_bytes(resp.content)
            out.append(f"{IMG_DIR.as_posix()}/{fname}")
        except Exception:
            continue
    return out


async def _read_upload(e):
    """Works with NiceGUI 1.x/2.x (e.content) and 3.x (e.file)."""
    if hasattr(e, "file"):
        return e.file.name, await e.file.read()
    return e.name, e.content.read()


def build_products_tab(card_css: str = "", button_css: str = ""):
    IMG_DIR.mkdir(parents=True, exist_ok=True)
    try:
        app.add_static_files(IMG_URL, str(IMG_DIR))
    except Exception:
        pass  # already registered

    data = _load()

    def save():
        ids = [p["id"] for p in data["products"]]
        if len(ids) != len(set(ids)) or "" in ids:
            ui.notify("Every product needs a unique ID", type="negative")
            return
        _save(data)
        ui.notify("Saved to data/products.json (product_tour.py picks it up on the next pass)", type="positive")

    def reload():
        data.clear()
        data.update(_load())
        product_list.refresh()

    async def add_product(raw: str = ""):
        info = {}
        raw = (raw or "").strip()
        if raw.startswith("http"):
            ui.notify("Reading the product in Chrome... if a check appears there, solve it.", timeout=8000)
            try:
                info = await run.io_bound(fetch_tiktok_product, raw)
            except Exception as ex:
                ui.notify(f"Browser problem: {ex}", type="negative", timeout=12000)
            if not info.get("name"):
                ui.notify("Couldn't read that page. Blank product created, fill it in manually.",
                          type="warning", timeout=8000)
        elif raw:
            try:
                info = json.loads(raw)
            except ValueError:
                ui.notify("Not a link or valid product data. Blank product created.", type="warning")

        n = len(data["products"]) + 1
        existing = {p["id"] for p in data["products"]}
        while f"P{n:03d}" in existing:
            n += 1
        pid = f"P{n:03d}"
        order = max([p["order"] for p in data["products"]] + [0]) + 1
        imgs = await run.io_bound(_download_images, pid, info.get("images", []))
        data["products"].append(_normalize({
            **NEW_PRODUCT_DEFAULTS,
            "id": pid, "order": order, "name": info.get("name") or "New product",
            "description": info.get("description", ""),
            "sizes_colors": info.get("sizes_colors", ""),
            "intro": info.get("intro", ""),
            "price": _money(info.get("price")),
            "original_price": _money(info.get("original_price")),
            "images": imgs,
        }))
        product_list.refresh()

    def delete_product(p):
        data["products"].remove(p)
        product_list.refresh()

    @ui.refreshable
    def product_list():
        for p in sorted(data["products"], key=lambda x: x["order"]):
            product_card(p)

    def product_card(p: dict):
        @ui.refreshable
        def gallery():
            with ui.row().classes("items-center gap-2"):
                for path in list(p["images"]):
                    with ui.card().tight().classes("w-28"):
                        ui.image(f"{IMG_URL}/{Path(path).name}").classes("h-28 object-cover")
                        ui.button(icon="delete", on_click=lambda _, path=path: remove_image(path)) \
                            .props("flat dense color=negative")
                if not p["images"]:
                    ui.label("No images yet").classes("text-grey")

        def remove_image(path):
            p["images"].remove(path)
            try:
                Path(path).unlink()
            except OSError:
                pass
            gallery.refresh()

        async def on_upload(e):
            name, content = await _read_upload(e)
            ext = Path(name).suffix.lower()
            if ext not in IMG_EXT:
                ui.notify(f"Unsupported file type: {ext}", type="negative")
                return
            fname = f"{p['id'] or 'product'}_{int(time.time() * 1000)}{ext}"
            (IMG_DIR / fname).write_bytes(content)
            p["images"].append(f"{IMG_DIR.as_posix()}/{fname}")
            gallery.refresh()

        with ui.expansion(f"#{p['order']}  {p['name']}  ({p['id']})", icon="shopping_bag").classes("w-full") \
                .props("group=products"):
            with ui.card().style(card_css).classes("w-full"):
                with ui.row():
                    ui.input("ID").bind_value(p, "id").style("width:100px;")
                    ui.number("Cart order", format="%d").bind_value(
                        p, "order", forward=lambda v: int(v or 0)).style("width:100px;")
                    ui.switch("Active (uncheck = sold out, skipped in tour)").bind_value(p, "active")
                for key, label in TEXT_FIELDS:
                    ui.input(label).bind_value(p, key).classes("w-full")
                    if key == "description":
                        ui.textarea("Introduction (one paragraph the AI reads when presenting this product)") \
                            .bind_value(p, "intro").props("autogrow outlined").classes("w-full")
                ui.input("Aliases (comma separated)").bind_value(
                    p, "aliases",
                    forward=lambda t: [x.strip() for x in (t or "").split(",") if x.strip()],
                    backward=lambda v: ", ".join(v or []),
                ).classes("w-full")
                ui.textarea("Highlights (one per line)").bind_value(
                    p, "highlights", forward=_from_lines, backward=_to_lines).classes("w-full")
                ui.textarea("FAQ (one per line:  question | answer)").bind_value(
                    p, "faq", forward=_faq_from_text, backward=_faq_to_text).classes("w-full")

                ui.label("Images")
                gallery()
                ui.upload(label="Add images", multiple=True, auto_upload=True, on_upload=on_upload) \
                    .props("accept=image/* flat bordered").classes("w-full")

                ui.button("Delete product", icon="delete", color="negative",
                          on_click=lambda _, p=p: delete_product(p)).props("flat")

    # ---- toolbar (must stay at this indent level, inside build_products_tab) ----
    with ui.dialog() as add_dialog, ui.card().classes("w-96"):
        link = ui.input("TikTok product link (empty = blank product)").classes("w-full")

        async def do_add():
            add_dialog.close()
            v = link.value or ""
            link.value = ""
            await add_product(v)

        ui.button("Add", on_click=do_add).style(button_css)

    with ui.row().classes("items-center"):
        ui.input("Shop name").bind_value(data, "shop_name").style("width:260px;")
        ui.button("Add", icon="add", on_click=add_dialog.open).style(button_css)
        ui.button("Save", icon="save", on_click=save).style(button_css)
        ui.button("Reload from file", icon="refresh", on_click=reload).props("flat")
    product_list()